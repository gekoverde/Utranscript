"""Regression tests for structured data: CSV typing and headers, INI sections, Excel cells,
non-text keys, bytes and the file names of workbook sheets."""

from __future__ import annotations

import base64
import configparser
import datetime
import json
import plistlib
import re
import zipfile

import pytest

from transcript.converters import data


def csv_text(path) -> str:
    return path.read_bytes().decode("utf-8-sig")  # without newline translation


def write_workbook(path, sheets: dict[str, list[list]]):
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


# reading ---------------------------------------------------------------------------------


def test_xlsx_with_a_wrong_dimension_keeps_every_row(samples, convert):
    good = write_workbook(samples / "good.xlsx", {"Hoja": [["a", "b", "c"], [1, 2, 3], [4, 5, 6, 7], [8]]})
    bad = samples / "bad.xlsx"
    with zipfile.ZipFile(good) as src, zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', content)
            dst.writestr(item, content)

    assert csv_text(convert(bad, "csv")) == "a,b,c,column4\r\n1,2,3,\r\n4,5,6,7\r\n8,,,\r\n"
    assert json.loads(convert(bad, "json").read_text(encoding="utf-8"))[2] == {"a": 8, "b": None, "c": None, "column4": None}


def test_ini_default_section_is_plain_and_overrides_are_kept(samples, convert):
    src = samples / "app.ini"
    src.write_text("[DEFAULT]\ntimeout = 30\n\n[server]\ntimeout = 60\nport = 8080\n", encoding="utf-8")
    result = json.loads(convert(src, "json").read_text(encoding="utf-8"))
    assert result == {"DEFAULT": {"timeout": "30"}, "server": {"timeout": "60", "port": "8080"}}


def test_xls_error_cells_become_their_error_text():
    xlrd = pytest.importorskip("xlrd")

    def value(code):
        return data._xls_value(xlrd.sheet.Cell(xlrd.XL_CELL_ERROR, code), 0)

    assert (value(7), value(42), value(15), value(99)) == ("#DIV/0!", "#N/A", "#VALUE!", "#ERR")


def test_csv_cells_beyond_the_header_are_kept(samples, convert):
    src = samples / "wide.csv"
    src.write_text("name,score\nalice,1,extra1\nbob,2,extra2\ncarol,3,\n", encoding="utf-8")
    assert json.loads(convert(src, "json").read_text(encoding="utf-8")) == [
        {"name": "alice", "score": 1, "column3": "extra1"},
        {"name": "bob", "score": 2, "column3": "extra2"},
        {"name": "carol", "score": 3, "column3": ""},
    ]
    trailing = samples / "trailing.csv"
    trailing.write_text("a,b\n1,2,\n3,4,,\n", encoding="utf-8")  # empty cells add no columns
    assert json.loads(convert(trailing, "json").read_text(encoding="utf-8")) == [{"a": 1, "b": 2}, {"a": 3, "b": 4}]


def test_csv_numbers(samples, convert):
    src = samples / "numbers.csv"
    src.write_text(
        "code,price,zero,big,mixed,ints\nA1,1.50,10,1.25,1,7\nB2,2.00,-0,12345678901234567.25,2.5,-8\n", encoding="utf-8"
    )
    # From CSV to TSV every cell stays exactly as written.
    assert csv_text(convert(src, "tsv")) == src.read_text(encoding="utf-8").replace(",", "\t").replace("\n", "\r\n")
    # Elsewhere a column is numeric when every value is exactly a number ('1.50' is 1.5), and
    # text when one is not ('-0', or more digits than a float can hold).
    assert json.loads(convert(src, "json").read_text(encoding="utf-8")) == [
        {"code": "A1", "price": 1.5, "zero": "10", "big": "1.25", "mixed": 1, "ints": 7},
        {"code": "B2", "price": 2.0, "zero": "-0", "big": "12345678901234567.25", "mixed": 2.5, "ints": -8},
    ]
    import openpyxl

    sheet = openpyxl.load_workbook(convert(src, "xlsx")).active
    assert [sheet.cell(row, 2).value for row in (2, 3)] == [1.5, 2.0]  # prices Excel can add up


def test_header_names_are_always_unique(samples, convert):
    assert data._unique_header(["a", "a", "a_2"]) == ["a", "a_3", "a_2"]
    assert data._unique_header(["", "column1", " ", None, "x", "x", "x"]) == [
        "column1_2",
        "column1",
        "column3",
        "column4",
        "x",
        "x_2",
        "x_3",
    ]
    src = samples / "dupes.csv"
    src.write_text("a,a,a_2\n1,2,3\n", encoding="utf-8")
    assert json.loads(convert(src, "json").read_text(encoding="utf-8")) == [{"a": 1, "a_3": 2, "a_2": 3}]


# writing ---------------------------------------------------------------------------------


@pytest.mark.parametrize("delimiter", [",", "\t"])
def test_header_only_table_keeps_its_header(samples, convert, delimiter):
    import openpyxl

    src = samples / ("empty.csv" if delimiter == "," else "empty.tsv")
    src.write_text(delimiter.join(["name", "email", "phone"]) + "\n", encoding="utf-8")
    other = "tsv" if delimiter == "," else "csv"
    assert csv_text(convert(src, other)) == ("\t" if other == "tsv" else ",").join(["name", "email", "phone"]) + "\r\n"
    workbook = openpyxl.load_workbook(convert(src, "xlsx"))
    assert [c.value for c in workbook.active[1]] == ["name", "email", "phone"]
    workbook.close()
    assert "<th>email</th>" in convert(src, "html").read_text(encoding="utf-8")


def test_header_only_sheet_in_a_workbook_keeps_its_header(samples, convert):
    src = write_workbook(samples / "book.xlsx", {"Datos": [["a", "b"], [1, 2]], "Vacia": [["x", "y"]]})
    folder = convert(src, "csv")
    assert csv_text(folder / "Vacia.csv") == "x,y\r\n"


def test_xlsx_text_is_never_a_formula(samples, convert):
    import openpyxl

    src = samples / "eq.csv"
    src.write_text('=title,b\n=== heading ===,=)\n"=HYPERLINK(""http://example.com"")",=\n', encoding="utf-8")
    result = convert(src, "xlsx")
    with zipfile.ZipFile(result) as zf:
        assert b"<f>" not in zf.read("xl/worksheets/sheet1.xml")
    workbook = openpyxl.load_workbook(result)
    values = [[c.value for c in row] for row in workbook.active.iter_rows()]
    workbook.close()
    assert values == [["=title", "b"], ["=== heading ===", "=)"], ['=HYPERLINK("http://example.com")', "="]]


def test_xlsx_drops_characters_it_cannot_store(samples, convert):
    import openpyxl

    src = samples / "bell.json"
    src.write_text(json.dumps([{"name\u0001": "bell\u0007x", "ok": "tab\tand\nnewline"}]), encoding="utf-8")
    workbook = openpyxl.load_workbook(convert(src, "xlsx"))
    assert [[c.value for c in row] for row in workbook.active.iter_rows()] == [["name", "ok"], ["bellx", "tab\tand\nnewline"]]
    workbook.close()


YAML_WITH_KEYS = """\
holidays:
  2024-01-01: New year
  2024-12-25 10:30:00: Christmas
  7: seven
  null: nothing
events:
  - 2024-05-01: may
    3: three
"""


@pytest.mark.parametrize("target", ["json", "jsonl", "toml", "xml", "plist", "ini", "csv", "tsv", "xlsx"])
def test_non_text_keys_are_written_as_text(samples, convert, target):
    pytest.importorskip("yaml")
    src = samples / "keys.yaml"
    src.write_text(YAML_WITH_KEYS, encoding="utf-8")
    assert convert(src, target).stat().st_size > 0


def test_non_text_keys_use_the_json_spelling(samples, convert):
    pytest.importorskip("yaml")
    src = samples / "keys.yaml"
    src.write_text(YAML_WITH_KEYS, encoding="utf-8")
    assert json.loads(convert(src, "json").read_text(encoding="utf-8")) == {
        "holidays": {"2024-01-01": "New year", "2024-12-25T10:30:00": "Christmas", "7": "seven", "null": "nothing"},
        "events": [{"2024-05-01": "may", "3": "three"}],
    }
    records = samples / "records.yaml"
    records.write_text("- 2024-01-01: a\n  1: b\n- 2024-01-01: c\n  1: d\n", encoding="utf-8")
    assert csv_text(convert(records, "csv")) == "2024-01-01,1\r\na,b\r\nc,d\r\n"


def test_bytes_become_base64_text(samples, convert):
    payload = bytes(range(0, 256, 17))
    src = samples / "blob.plist"
    with open(src, "wb") as fh:
        plistlib.dump({"payload": payload, "when": datetime.datetime(2024, 1, 2, 3, 4, 5)}, fh)
    result = json.loads(convert(src, "json").read_text(encoding="utf-8"))
    assert base64.b64decode(result["payload"]) == payload
    assert result["when"] == "2024-01-02T03:04:05"
    assert base64.b64encode(payload).decode("ascii") in csv_text(convert(src, "csv"))


def test_records_are_written_as_numbered_ini_sections(samples, convert):
    src = samples / "people.csv"
    src.write_text("name,age\nAna,30\nLuis,\n", encoding="utf-8")
    parser = configparser.ConfigParser(interpolation=None)
    parser.read_string(convert(src, "ini").read_text(encoding="utf-8"))
    assert {s: dict(parser[s]) for s in parser.sections()} == {
        "1": {"name": "Ana", "age": "30"},
        "2": {"name": "Luis", "age": ""},
    }
    lines = samples / "rows.jsonl"
    lines.write_text('{"a": 1}\n{"a": 2, "b": [1, 2]}\n', encoding="utf-8")
    assert "[2]\na = 2\nb = [1, 2]" in convert(lines, "ini").read_text(encoding="utf-8")


def test_sheet_file_names_are_unique_and_portable(samples, convert):
    sheets = ["Total<2024", "Total>2024", "Aux", "total_2024", "Trail. ", "nul.backup", "COM1"]
    src = write_workbook(samples / "libro.xlsx", {name: [["n"], [i]] for i, name in enumerate(sheets)})
    folder = convert(src, "csv")
    names = {"Total_2024", "Total_2024 (2)", "_Aux", "total_2024 (3)", "Trail", "_nul.backup", "_COM1"}
    assert sorted(p.name for p in folder.iterdir()) == sorted(n + ".csv" for n in names)
    assert csv_text(folder / "Total_2024 (2).csv") == "n\r\n1\r\n"  # the second sheet was not overwritten


def test_sheet_file_stem_rules():
    used: set[str] = set()
    stems = [data._file_stem(name, used) for name in ["CON", "con", "LPT9.x", "Résumé", "résumé", "...", "a?b"]]
    assert stems == ["_CON", "_con (2)", "_LPT9.x", "Résumé", "résumé (2)", "sheet", "a_b"]
